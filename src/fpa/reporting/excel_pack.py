"""Monthly variance pack for budget owners (Excel): a summary tab, one tab per department and
the full-year outlook."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1F3A5F"
FILLS = {"Red": "F8D7DA", "Amber": "FFE8B3", "Green": "D9EAD3"}
USD = '#,##0;[Red](#,##0)'
PCT = '0.0%;[Red]-0.0%'
THIN = Border(bottom=Side(style="thin", color="D0D7DE"))


def _header(ws, row: int, headers: list[str], widths: list[int]) -> None:
    for c, (h, w) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[row].height = 30


def _title(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="555555")


def _write_rows(ws, start: int, rows: list[list], formats: list[str | None], status_col: int | None = None, wrap_col: int | None = None) -> int:
    r = start
    for values in rows:
        for c, (v, fmt) in enumerate(zip(values, formats, strict=True), start=1):
            cell = ws.cell(row=r, column=c, value=None if (isinstance(v, float) and pd.isna(v)) else v)
            if fmt:
                cell.number_format = fmt
            cell.border = THIN
            cell.alignment = Alignment(vertical="top", wrap_text=(c == wrap_col))
        if status_col:
            s = values[status_col - 1]
            if s in FILLS:
                ws.cell(row=r, column=status_col).fill = PatternFill("solid", fgColor=FILLS[s])
        r += 1
    return r


def write_variance_pack(marts_dir: Path, as_of: pd.Period, out_dir: Path, company: str) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    flags = pd.read_parquet(marts_dir / "variance_monthly.parquet")
    ytd = pd.read_parquet(marts_dir / "variance_ytd.parquet")
    outlook = pd.read_parquet(marts_dir / "fy_outlook.parquet")
    cur = flags[flags["period"] == str(as_of)].copy()
    order = {"Red": 0, "Amber": 1, "Green": 2}
    cur = cur.sort_values(["severity", "impact_usd"], key=lambda s: s.map(order) if s.name == "severity" else s)
    month = as_of.strftime("%B %Y")

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _title(ws, f"{company} - Monthly Variance Pack", f"{month} close | Budget = AOP FY{as_of.year} | Amounts in USD")
    counts = cur["severity"].value_counts()
    ws["A4"], ws["B4"] = "Red - investigate", int(counts.get("Red", 0))
    ws["A5"], ws["B5"] = "Amber - explain / monitor", int(counts.get("Amber", 0))
    ws["A6"], ws["B6"] = "Green - on track", int(counts.get("Green", 0))
    for r, s in ((4, "Red"), (5, "Amber"), (6, "Green")):
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=FILLS[s])
    headers = ["Status", "Department", "P&L line", "Actual", "Budget", "Var $", "Var %", "Last forecast", "Expected low", "Expected high", "Commentary (draft)"]
    widths = [9, 20, 22, 13, 13, 12, 9, 13, 13, 13, 90]
    _header(ws, 8, headers, widths)
    needs = cur[cur["severity"] != "Green"]
    rows = needs[["severity", "department", "fs_line", "actual_usd", "budget_usd", "var_budget_usd", "var_budget_pct", "forecast_usd", "range_lo_usd", "range_hi_usd", "commentary"]].values.tolist()
    _write_rows(ws, 9, rows, [None, None, None, USD, USD, USD, PCT, USD, USD, USD, None], status_col=1, wrap_col=11)
    ws.freeze_panes = "A9"

    ytd_idx = ytd.set_index("unique_id")
    for dept, g in cur.groupby("department", sort=True):
        sh = wb.create_sheet(dept[:31])
        _title(sh, f"{dept} - {month}", "Month and year-to-date versus budget")
        headers = ["Status", "P&L line", "Actual", "Budget", "Var $", "Var %", "Prior year", "YTD actual", "YTD budget", "YTD var $", "YTD var %", "Commentary (draft)"]
        _header(sh, 4, headers, [9, 24, 13, 13, 12, 9, 13, 14, 14, 13, 10, 80])
        rows = []
        for r in g.itertuples():
            y = ytd_idx.loc[r.unique_id] if r.unique_id in ytd_idx.index else None
            rows.append([
                r.severity, r.fs_line, r.actual_usd, r.budget_usd, r.var_budget_usd, r.var_budget_pct, r.prior_year_usd,
                None if y is None else y["ytd_actual_usd"], None if y is None else y["ytd_budget_usd"],
                None if y is None else y["var_usd"], None if y is None else y["var_pct"], r.commentary,
            ])
        _write_rows(sh, 5, rows, [None, None, USD, USD, USD, PCT, USD, USD, USD, USD, PCT, None], status_col=1, wrap_col=12)
        sh.freeze_panes = "A5"

    sh = wb.create_sheet("FY Outlook")
    fy, n_act = int(outlook["fiscal_year"].iloc[0]), int(outlook["months_actual"].iloc[0])
    _title(sh, f"FY{fy} landing estimate", f"{n_act} months of actuals + rolling forecast for the remaining {12 - n_act}")
    headers = ["Department", "P&L line", "YTD actual", "Remaining forecast", "FY outlook", "FY budget", "Outlook vs budget"]
    _header(sh, 4, headers, [20, 24, 14, 16, 14, 14, 16])
    o = outlook.sort_values(["account_type", "fs_line", "department"], ascending=[False, True, True])
    rows = o[["department", "fs_line", "ytd_actual_usd", "remaining_forecast_usd", "fy_outlook_usd", "fy_budget_usd", "outlook_vs_budget_usd"]].values.tolist()
    _write_rows(sh, 5, rows, [None, None, USD, USD, USD, USD, USD])
    sh.freeze_panes = "A5"

    path = out_dir / f"variance_pack_{as_of}.xlsx"
    for old in out_dir.glob("variance_pack_*.xlsx"):
        if old != path:
            old.unlink()
    wb.save(path)
    return path
